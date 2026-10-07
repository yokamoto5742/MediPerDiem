from pathlib import Path

import openpyxl

from service.average_workbook import write_average_workbook


def make_trend_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "外来日当円"
    sheet.append(["ID", "医師名", "診療科", 202607, 202608])
    sheet.append([103, "内科 太郎", "内", 11245.85, 10639.38])
    sheet.append([107, "眼科 花子", "眼", None, 0])
    sheet.append([144, "麻酔 三郎", "麻", 3304.5, 999.4])
    workbook.save(path)
    return path


def test_writes_recipient_rows_with_sum_and_average(tmp_path: Path) -> None:
    trend_path = make_trend_workbook(tmp_path / "trend.xlsx")
    recipients_path = tmp_path / "recipients.csv"
    recipients_path.write_text(
        "ID,メールアドレス\n107,花子 <h@example.com>\n999,不明 <x@example.com>\n103,太郎 <t@example.com>\n",
        encoding="utf-8-sig",
    )
    output_path = tmp_path / "average.xlsx"

    write_average_workbook(trend_path, recipients_path, output_path)

    sheet = openpyxl.load_workbook(output_path)["一覧"]
    assert [[cell.value for cell in row] for row in sheet.iter_rows()] == [
        ["ID", "医師名", "メールアドレス", 202607, 202608],
        [107, "眼科 花子", None, None, 0],
        [103, "内科 太郎", None, 11246, 10639],
        [None, "外来日当円合計", None, "=SUM(D2:D3)", "=SUM(E2:E3)"],
        [None, "外来日当円平均", None, "=AVERAGE(D2:D3)", "=AVERAGE(E2:E3)"],
    ]
    assert sheet["D3"].number_format == sheet["E5"].number_format == "#,##0"


def test_overwrites_existing_workbook(tmp_path: Path) -> None:
    trend_path = make_trend_workbook(tmp_path / "trend.xlsx")
    recipients_path = tmp_path / "recipients.csv"
    recipients_path.write_text("ID,メールアドレス\n144,三郎 <s@example.com>\n", encoding="utf-8-sig")
    output_path = tmp_path / "average.xlsx"
    write_average_workbook(trend_path, recipients_path, output_path)

    recipients_path.write_text("ID,メールアドレス\n103,太郎 <t@example.com>\n", encoding="utf-8-sig")
    write_average_workbook(trend_path, recipients_path, output_path)

    sheet = openpyxl.load_workbook(output_path)["一覧"]
    assert [sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)] == ["ID", 103, None, None]
