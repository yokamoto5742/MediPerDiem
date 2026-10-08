from pathlib import Path

import openpyxl

from service.revenue_reader import DoctorRevenue
from service.trend_workbook import backup_workbook, update_trend_workbook

NUMBER_FORMAT = "#,##0_ "


def make_trend_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    default_sheet = workbook.active
    assert default_sheet is not None
    workbook.remove(default_sheet)
    for title in ("外来日当円", "外来収益合計"):
        sheet = workbook.create_sheet(title)
        sheet.append(["ID", "医師名", "診療科", 202606, 202607])
        sheet.append([103, "内科 太郎", "内", 100.5, 200.5])
        sheet.append([107, "眼科 花子", "眼", 300, None])
        for row in (2, 3):
            sheet.cell(row, 5).number_format = NUMBER_FORMAT
        sheet.column_dimensions["E"].width = 11.5
    workbook.save(path)
    return path


def test_appends_month_column_to_both_sheets(tmp_path: Path) -> None:
    path = make_trend_workbook(tmp_path / "trend.xlsx")
    revenues = [DoctorRevenue(103, "内科 太郎", 4096162, 10639.38), DoctorRevenue(107, "眼科 花子", 21240, None)]

    assert update_trend_workbook(path, 202608, revenues) == []

    workbook = openpyxl.load_workbook(path)
    per_diem, revenue = workbook["外来日当円"], workbook["外来収益合計"]
    assert [per_diem.cell(row, 6).value for row in (1, 2, 3)] == [202608, 10639.38, None]
    assert [revenue.cell(row, 6).value for row in (1, 2, 3)] == [202608, 4096162, 21240]
    assert per_diem.cell(2, 5).value == 200.5
    assert per_diem.cell(2, 6).number_format == NUMBER_FORMAT
    assert per_diem.column_dimensions["F"].width == 11.5


def test_rerun_overwrites_same_month_column(tmp_path: Path) -> None:
    path = make_trend_workbook(tmp_path / "trend.xlsx")
    update_trend_workbook(path, 202608, [DoctorRevenue(103, "", 100, 10), DoctorRevenue(107, "", 200, 20)])
    update_trend_workbook(path, 202608, [DoctorRevenue(103, "", 300, 30)])

    sheet = openpyxl.load_workbook(path)["外来日当円"]
    assert [sheet.cell(row, 6).value for row in (1, 2, 3)] == [202608, 30, None]
    assert sheet.cell(1, 7).value is None


def test_inserts_missing_doctor_below_same_department(tmp_path: Path) -> None:
    path = make_trend_workbook(tmp_path / "trend.xlsx")
    new_doctor = DoctorRevenue(440, "内科 次郎", 116821, 5310.05, "内")

    assert update_trend_workbook(path, 202608, [new_doctor, DoctorRevenue(107, "", 200, 20, "眼")]) == [
        (3, new_doctor)
    ]

    workbook = openpyxl.load_workbook(path)
    per_diem, revenue = workbook["外来日当円"], workbook["外来収益合計"]
    assert [cell.value for cell in per_diem[3]] == [440, "内科 次郎", "内", None, None, 5310.05]
    assert [cell.value for cell in revenue[3]] == [440, "内科 次郎", "内", None, None, 116821]
    assert per_diem.cell(3, 5).number_format == NUMBER_FORMAT
    assert [per_diem.cell(4, column).value for column in (1, 5, 6)] == [107, None, 20]


def test_inserts_doctor_of_new_department_at_bottom(tmp_path: Path) -> None:
    path = make_trend_workbook(tmp_path / "trend.xlsx")
    new_doctor = DoctorRevenue(899, "泌尿 四郎", 300, 30, "泌")

    assert update_trend_workbook(path, 202608, [new_doctor]) == [(4, new_doctor)]

    sheet = openpyxl.load_workbook(path)["外来日当円"]
    assert [cell.value for cell in sheet[4]] == [899, "泌尿 四郎", "泌", None, None, 30]


def test_returns_row_numbers_after_all_insertions(tmp_path: Path) -> None:
    path = make_trend_workbook(tmp_path / "trend.xlsx")
    eye_doctor = DoctorRevenue(500, "眼科 三郎", 300, 30, "眼")
    internal_doctor = DoctorRevenue(440, "内科 次郎", 100, 10, "内")

    # 後から挿入した内科の行により、先に挿入した眼科の行は1行下へずれる
    assert update_trend_workbook(path, 202608, [eye_doctor, internal_doctor]) == [
        (5, eye_doctor), (3, internal_doctor)
    ]


def test_rerun_does_not_insert_doctor_twice(tmp_path: Path) -> None:
    path = make_trend_workbook(tmp_path / "trend.xlsx")
    new_doctor = DoctorRevenue(440, "内科 次郎", 100, 10, "内")
    update_trend_workbook(path, 202608, [new_doctor])

    assert update_trend_workbook(path, 202608, [new_doctor]) == []
    assert openpyxl.load_workbook(path)["外来日当円"].max_row == 4


def test_backup_keeps_original_content(tmp_path: Path) -> None:
    path = make_trend_workbook(tmp_path / "trend.xlsx")
    backup_path = backup_workbook(path, 12)
    update_trend_workbook(path, 202608, [DoctorRevenue(103, "", 100, 10)])

    assert backup_path.parent == path.parent / "backup"
    assert openpyxl.load_workbook(backup_path)["外来日当円"].cell(1, 6).value is None


def test_backup_keeps_only_latest_generations(tmp_path: Path) -> None:
    path = make_trend_workbook(tmp_path / "trend.xlsx")
    backup_dir = tmp_path / "backup"
    backup_dir.mkdir()
    for day in (1, 2, 3):
        (backup_dir / f"trend_2025040{day}_000000.xlsx").touch()
    unrelated = backup_dir / "other_20250401_000000.xlsx"
    unrelated.touch()

    backup_path = backup_workbook(path, 2)

    assert sorted(backup_dir.glob("trend_*")) == [backup_dir / "trend_20250403_000000.xlsx", backup_path]
    assert unrelated.exists()
