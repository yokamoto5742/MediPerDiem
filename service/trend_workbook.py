import glob
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from service.revenue_reader import DoctorRevenue

PER_DIEM_SHEET_NAME = "外来日当円"
REVENUE_SHEET_NAME = "外来収益合計"
DEPARTMENT_COLUMN = 3
FIRST_MONTH_COLUMN = 4
BACKUP_DIR_NAME = "backup"


def backup_workbook(trend_path: Path, generations: int) -> Path:
    """backupフォルダへ更新前の変化表を複製し、最新generations世代だけ残す"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = trend_path.parent / BACKUP_DIR_NAME
    backup_dir.mkdir(exist_ok=True)
    backup_path = backup_dir / f"{trend_path.stem}_{timestamp}{trend_path.suffix}"
    shutil.copy2(trend_path, backup_path)

    # ファイル名のタイムスタンプ順 = 時系列順
    backups = sorted(backup_dir.glob(f"{glob.escape(trend_path.stem)}_*{trend_path.suffix}"))
    for old_backup in backups[:max(len(backups) - generations, 0)]:
        old_backup.unlink()
    return backup_path


def update_trend_workbook(
    trend_path: Path, month: int, revenues: list[DoctorRevenue]
) -> list[tuple[int, DoctorRevenue]]:
    """対象月の列を両シートに書き込み、変化表に行が無く新たに挿入した(外来日当円シートの行番号, 医師)を返す"""
    workbook = openpyxl.load_workbook(trend_path)
    per_diem_sheet = workbook[PER_DIEM_SHEET_NAME]
    revenue_sheet = workbook[REVENUE_SHEET_NAME]
    registered_ids = _write_month_column(per_diem_sheet, month, {r.doctor_id: r.per_diem for r in revenues})
    _write_month_column(revenue_sheet, month, {r.doctor_id: r.revenue for r in revenues})
    unregistered = [r for r in revenues if r.doctor_id not in registered_ids]
    for doctor in unregistered:
        _insert_doctor_row(per_diem_sheet, month, doctor, doctor.per_diem)
        _insert_doctor_row(revenue_sheet, month, doctor, doctor.revenue)
    workbook.save(trend_path)
    # 後続の挿入で行がずれるため、行番号はすべて挿入し終えてから求める
    row_by_id = {per_diem_sheet.cell(row, 1).value: row for row in range(2, per_diem_sheet.max_row + 1)}
    return [(row_by_id[doctor.doctor_id], doctor) for doctor in unregistered]


def _insert_doctor_row(sheet: Worksheet, month: int, doctor: DoctorRevenue, value: float | None) -> None:
    """同じ診療科の最終行の直下(診療科の行が無ければ表の末尾)に医師の行を挿入する"""
    doctor_rows = [row for row in range(2, sheet.max_row + 1) if isinstance(sheet.cell(row, 1).value, int)]
    department_rows = [
        row for row in doctor_rows if sheet.cell(row, DEPARTMENT_COLUMN).value == doctor.department
    ]
    above_row = (department_rows or doctor_rows)[-1]
    new_row = above_row + 1
    sheet.insert_rows(new_row)
    for column in range(1, sheet.max_column + 1):
        sheet.cell(new_row, column)._style = copy(sheet.cell(above_row, column)._style)
    sheet.cell(new_row, 1, doctor.doctor_id)
    sheet.cell(new_row, 2, doctor.name)
    sheet.cell(new_row, DEPARTMENT_COLUMN, doctor.department)
    sheet.cell(new_row, _find_month_column(sheet, month), value)


def _write_month_column(sheet: Worksheet, month: int, value_by_id: dict[int, float | None]) -> set[int]:
    column = _find_month_column(sheet, month)
    letter = get_column_letter(column)
    if sheet.cell(1, column).value is None:
        _copy_column_format(sheet, column - 1, column)
        sheet[f"{letter}1"] = month

    registered_ids: set[int] = set()
    for row in range(2, sheet.max_row + 1):
        doctor_id = sheet.cell(row, 1).value
        if isinstance(doctor_id, int):
            registered_ids.add(doctor_id)
            # 入力に無い医師は空欄にする(同じ月の再実行で前回の値を残さない)
            sheet[f"{letter}{row}"] = value_by_id.get(doctor_id)
    return registered_ids


def _find_month_column(sheet: Worksheet, month: int) -> int:
    """対象月の列を返す。無ければ見出し行の次の空き列を返す"""
    column = FIRST_MONTH_COLUMN
    while sheet.cell(1, column).value not in (None, month):
        column += 1
    return column


def _copy_column_format(sheet: Worksheet, source_column: int, target_column: int) -> None:
    for row in range(1, sheet.max_row + 1):
        sheet.cell(row, target_column)._style = copy(sheet.cell(row, source_column)._style)
    source_width = sheet.column_dimensions[get_column_letter(source_column)].width
    sheet.column_dimensions[get_column_letter(target_column)].width = source_width
