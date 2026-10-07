import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

LIST_SHEET_NAME = "一覧"
HEADER_ROW = 3
ADJUSTED_DEPARTMENT = "内"
EXCLUDED_DEPARTMENT = "健"
# 一覧シートの列位置(0始まり)と、その列にあるべき見出し
COL_DEPARTMENT, COL_ID, COL_NAME, COL_MEDICINE, COL_INJECTION, COL_TOTAL, COL_PATIENTS = 0, 1, 2, 4, 5, 13, 14
EXPECTED_HEADERS = {
    COL_ID: "ID",
    COL_MEDICINE: "クスリ",
    COL_INJECTION: "注射",
    COL_TOTAL: "合計",
    COL_PATIENTS: "外来人数合計",
}


@dataclass(frozen=True)
class DoctorRevenue:
    doctor_id: int
    name: str
    revenue: float
    per_diem: float | None


def parse_target_month(revenue_path: Path) -> int:
    """ファイル名末尾のYYYYMMを対象月として返す"""
    match = re.search(r"(\d{4})(0[1-9]|1[0-2])$", revenue_path.stem)
    if match is None:
        raise ValueError(f"ファイル名から対象月(YYYYMM)を取得できません: {revenue_path.name}")
    return int(match.group(0))


def read_doctor_revenues(revenue_path: Path) -> list[DoctorRevenue]:
    workbook = openpyxl.load_workbook(revenue_path, read_only=True, data_only=True)
    try:
        if LIST_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"「{LIST_SHEET_NAME}」シートがありません: {revenue_path.name}")
        rows = list(workbook[LIST_SHEET_NAME].iter_rows(min_row=HEADER_ROW, max_col=COL_PATIENTS + 1, values_only=True))
    finally:
        workbook.close()

    for column, expected in EXPECTED_HEADERS.items():
        if not rows or rows[0][column] != expected:
            raise ValueError(f"「{LIST_SHEET_NAME}」シートの列構成が想定と異なります(見出し「{expected}」)")

    # 「◆マスタなし」や合計行はIDが無いので除外される
    doctor_rows = [
        row for row in rows[1:]
        if isinstance(row[COL_ID], int) and row[COL_DEPARTMENT] != EXCLUDED_DEPARTMENT
    ]
    rows_by_name: dict[str, list[tuple]] = {}
    for row in doctor_rows:
        rows_by_name.setdefault(_name_key(row), []).append(row)
    return [_to_doctor_revenue(same_name_rows) for same_name_rows in rows_by_name.values()]


def _name_key(row: tuple) -> str:
    """同姓同名の判定キー。空白(全角・半角)の違いは無視する"""
    return re.sub(r"\s", "", str(row[COL_NAME] or "")) or str(row[COL_ID])


def _to_doctor_revenue(same_name_rows: list[tuple]) -> DoctorRevenue:
    """同姓同名の行を、IDが最も若い医師に合算する"""
    youngest = min(same_name_rows, key=lambda row: row[COL_ID])
    revenue = sum(_adjusted_revenue(row) for row in same_name_rows)
    patients = sum(row[COL_PATIENTS] or 0 for row in same_name_rows)
    return DoctorRevenue(
        doctor_id=youngest[COL_ID],
        name=str(youngest[COL_NAME] or ""),
        revenue=revenue,
        per_diem=revenue / patients if patients else None,
    )


def _adjusted_revenue(row: tuple) -> float:
    revenue = row[COL_TOTAL] or 0
    if row[COL_DEPARTMENT] == ADJUSTED_DEPARTMENT:
        # 内科はクスリと注射を1/10に調整する(「内科 (調整済み）」シートの合計と同じ)
        medicine_and_injection = (row[COL_MEDICINE] or 0) + (row[COL_INJECTION] or 0)
        revenue = revenue - medicine_and_injection + medicine_and_injection / 10
        if revenue == int(revenue):
            revenue = int(revenue)
    return revenue
