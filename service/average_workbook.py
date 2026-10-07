from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from service.distribution_csv import CSV_FIXED_HEADERS, read_distribution_rows, round_per_diem

SHEET_NAME = "一覧"
SUMMARY_ROWS = (("外来日当円合計", "SUM"), ("外来日当円平均", "AVERAGE"))
NUMBER_FORMAT = "#,##0"
NAME_COLUMN_WIDTH = 18


def write_average_workbook(trend_path: Path, recipients_path: Path, output_path: Path) -> None:
    """配信用CSVと同じ医師・月の表の下に、各月の合計と平均の行を加えたブックを作り直す"""
    months, doctor_rows, _ = read_distribution_rows(trend_path, recipients_path)
    workbook = openpyxl.Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = SHEET_NAME
    sheet.append(CSV_FIXED_HEADERS + months)
    for doctor_id, name, _, per_diems in doctor_rows:
        # メールアドレス列は見出しだけ残して空欄にする
        sheet.append([doctor_id, name, None] + [round_per_diem(value) for value in per_diems])

    last_doctor_row = sheet.max_row
    first_month_column = len(CSV_FIXED_HEADERS) + 1
    month_letters = [get_column_letter(first_month_column + index) for index in range(len(months))]
    for label, function in SUMMARY_ROWS:
        formulas = [f"={function}({letter}2:{letter}{last_doctor_row})" for letter in month_letters]
        sheet.append([None, label, None] + formulas)

    for row in sheet.iter_rows(min_row=2, min_col=first_month_column):
        for cell in row:
            cell.number_format = NUMBER_FORMAT
    sheet.column_dimensions["B"].width = NAME_COLUMN_WIDTH
    sheet.freeze_panes = "C2"
    workbook.save(output_path)
